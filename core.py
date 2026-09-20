import base64
import copy
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import sqlite3
import sys
import threading
from datetime import datetime, timezone
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
import unicodedata

from model import SKILL, config_status, read_config, messages_for, call_model, audit_brief, compact_brief, ModelError

ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(SKILL/'scripts'))
import validate_brief as validator
import build_brief as word_builder
import build_input_form as excel_builder

def now(): return datetime.now(timezone.utc).isoformat()
def uid(): return uuid4().hex
def clone(v): return copy.deepcopy(v)

class AppError(Exception):
    def __init__(self,message,status=400): self.status=status;super().__init__(message)

def checked(d):
    if not isinstance(d,dict): return ['任务书数据应为对象']
    if not isinstance(d.get('inputs',[]),list) or any(not isinstance(x,dict) for x in d.get('inputs',[])):
        return ['inputs须为输入对象列表']
    if any(not isinstance(d.get(k,{}),dict) for k in ['task_fit','scope_assessment','sample']):
        return ['task_fit、scope_assessment和sample须为对象']
    # Extend input responsibility without modifying the pinned Skill package.
    legacy=clone(d)
    student_inputs=[x for x in d.get('inputs',[]) if isinstance(x,dict) and x.get('owner')=='student']
    if student_inputs:
        for x in legacy.get('inputs',[]):
            if x.get('owner')=='student': x['owner']='customer' if x.get('required') else 'customer_optional'
    if d.get('task_fit',{}).get('private_data_provider')=='student' and student_inputs and not d.get('scope_assessment',{}).get('enterprise_dependency'):
        legacy['task_fit']['private_data_required']=False
    try: errors=validator.validate(legacy)
    except (TypeError,AttributeError,KeyError,ValueError): return ['任务书字段类型不符合 v0.4.3 约定，请检查数据结构。']
    for x in student_inputs:
        if x.get('dependency_note'): errors.append('团队自行准备的输入不能标为企业依赖；企业依赖须单独如实记录')
        if x.get('student_fallback'): errors.append('团队负责准备不是企业未提供时的替代项，请在requirement写清责任')
    if d.get('scope_assessment',{}).get('enterprise_dependency') and not any(x.get('owner')=='customer' and x.get('required') and x.get('dependency_note') for x in d.get('inputs',[])):
        errors.append('企业依赖任务必须有真实企业必填输入，团队输入不能替代')
    minimum=d.get('sample',{}).get('minimum_count')
    if minimum is not None:
        if type(minimum) is not int or minimum<1: errors.append('sample.minimum_count须为正整数')
        elif type(d.get('sample',{}).get('count',3)) is int and d.get('sample',{}).get('count',3)<min(3,minimum): errors.append('样例填写列数不足以承载原文最少样例数')
    if '输入不得索取或收集凭据' in errors:
        statements=[s for i in d.get('inputs',[]) for s in re.split(r'[。；;\n]',i.get('name','')+'。'+i.get('requirement','')) if re.search(r'密码|Cookie|cookie|验证码|密钥|登录态|访问令牌',s)]
        if all(re.search(r'不提供|不收集|不保存|不含|无需|不要|禁止|不得|不提交|不填写|不出现|不可接入|不可提供|不在.{0,12}填写',s) for s in statements):
            errors=[e for e in errors if e!='输入不得索取或收集凭据']
    for k,limit in [('scenario',130),('purpose',150),('objective',110)]:
        if f'{k}过长或类型错误' in errors:
            errors.remove(f'{k}过长或类型错误')
            errors.append(f'{k}必须是最多{limit}字的字符串，标点计数；详情移到输入、输出或规则字段，不得遗漏硬要求。')
    rules=d.get('rules',[])
    if not isinstance(rules,list): return errors+['处理规则必须是列表']
    seen=set()
    for r in rules:
        if not isinstance(r,dict) or any(not isinstance(r.get(k),str) or not r[k].strip() for k in ['id','name','detail','source']):
            errors.append('每条处理规则需有编号、名称、内容和来源');continue
        if r['id'] in seen: errors.append('处理规则编号重复')
        seen.add(r['id'])
        if r['source'] not in d.get('sources',{}) and r['source'] not in ['proposal','public-practice']: errors.append('处理规则来源不存在：'+r['id'])
    return errors

def sample_request(d):
    sample=d.get('sample',{})
    minimum=sample.get('minimum_count')
    count=sample.get('count',3)
    amount=f'至少{minimum}组' if minimum else (f'{count}组' if count<3 else '1至3组')
    customer=sum(x.get('required') is True and x.get('owner')=='customer' for x in d.get('inputs',[]))
    student=sum(x.get('required') is True and x.get('owner')=='student' for x in d.get('inputs',[]))
    if student:
        responsibility=(f'企业提供{customer}项、' if customer else '无需企业代为准备；')+f'团队自行准备{student}项必填输入。'
        return responsibility+f'请按输入表的责任分工准备{amount}样例并记录文件位置。'
    return f'请按上方“企业提供”的{customer}项必填内容，提供{amount}示例填写内容。'

def delivery_data(d):
    data=clone(d)
    for item in data.get('inputs',[]):
        if item.get('owner')=='student':
            item['requirement']='由团队自行准备。'+item['requirement']
    return data

def confirmation_rows(d):
    try: rows=validator.confirmations(d)
    except (TypeError,KeyError,ValueError,AttributeError): return []
    hidden=set(d.get('dismissed_confirmations',[]))
    rows=[r for r in rows if r['id'] not in hidden]
    rows+=clone(d.get('custom_confirmations',[]))
    answers=d.get('confirmation_answers',{})
    for row in rows:
        if row['id']=='C-SAMPLE':
            row['question']=sample_request(d)
            if d.get('sample',{}).get('real'): row['question']+='已有真实样例，请核对覆盖，不足时补充。'
        a=answers.get(row['id'],{})
        if a.get('text','').strip(): row['answer']=a['text']
        if row['id']=='C-RUBRIC' and d.get('rubric_approved'):
            row['answer']='☑ 确认当前评分规则'+ ('\n'+a['text'] if a.get('text','').strip() else '')
        row['recorded']=bool(a.get('text','').strip()) or (row['id']=='C-FORMAT' and bool(d.get('selected_output'))) or (row['id']=='C-RUBRIC' and bool(d.get('rubric_approved')))
        row['by']=a.get('by','');row['at']=a.get('at','')
    return rows

def preserve_confirmation_controls(old,new):
    for key in ('custom_confirmations','dismissed_confirmations'):
        new[key]=clone((old or {}).get(key,[]))

def rubric_changed(old,new):
    return bool(old) and any(new.get(k)!=old.get(k) for k in ('outputs','kr2_2','kr2_3'))

def rubric_data():
    return {'kr1_dimensions':word_builder.KR1_DIMENSIONS,'quality_scale':word_builder.QUALITY_SCALE,'usage_scale':word_builder.USAGE_SCALE}

# Adapter uses the pinned generator unchanged; answers remain a single data source.
word_builder.confirmations=confirmation_rows

def extract_file(path,name):
    ext=Path(name).suffix.lower()
    if ext in ['.md','.txt','.csv']:
        try: text=path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError:
            try: text=path.read_text(encoding='gb18030')
            except UnicodeDecodeError: raise AppError('文字编码无法识别，请另存为UTF-8。')
    elif ext in ['.docx','.xlsx']:
        try:
            with ZipFile(path) as z:
                if len(z.infolist())>3000 or sum(i.file_size for i in z.infolist())>60_000_000: raise AppError('文件解压后过大，请拆分文件。')
        except BadZipFile: raise AppError('Office文件损坏或格式与扩展名不符。')
        if ext=='.docx':
            from docx import Document
            from docx.oxml.ns import qn
            doc=Document(path)
            text='\n'.join(''.join(e.itertext()) if False else ''.join(t.text or '' for t in e.iter(qn('w:t'))) for e in doc.element.body.iter(qn('w:p')))
        else:
            from openpyxl import load_workbook
            wb=load_workbook(path,read_only=True,data_only=False)
            pieces=[];total=0
            for sh in wb:
                pieces.append('工作表：'+sh.title)
                for i,row in enumerate(sh.iter_rows(values_only=True),1):
                    total+=1
                    if total>10000: wb.close();raise AppError('表格超过10000行，请只上传与需求有关的区域。')
                    if any(v is not None for v in row): pieces.append(f'行{i}：'+' | '.join(str(v) if v is not None else '' for v in row))
            wb.close();text='\n'.join(pieces)
    elif ext=='.pdf':
        from pypdf import PdfReader
        doc=PdfReader(path)
        if doc.is_encrypted: raise AppError('PDF已加密，请提供可读取版本。')
        if len(doc.pages)>100: raise AppError('PDF超过100页，请拆分需求材料。')
        pieces=[]
        for i,p in enumerate(doc.pages,1):
            s=p.extract_text() or ''
            if len(s.strip())<8: raise AppError(f'PDF第{i}页没有足够可读取文字，可能是扫描件。请补充该页文字；本版不静默跳过。')
            pieces.append(f'第{i}页\n{s}')
        text='\n'.join(pieces)
    else: raise AppError('支持 MD、TXT、DOCX、PDF、XLSX、CSV 文件。')
    text=unicodedata.normalize('NFKC',text)
    text=''.join(c for c in text if c in '\n\t' or ord(c)>=32)
    if not text.strip(): raise AppError('没有提取到可读取的正文。')
    if len(text)>120000: raise AppError('单文件正文超过12万字，请拆分材料。')
    return '\n'.join(f'L{i}: {line}' for i,line in enumerate(text.splitlines(),1))

class Store:
    def __init__(self,data=None):
        self.root=Path(data or os.environ.get('BRIEF_DATA_DIR',ROOT/'data')).resolve()
        self.root.mkdir(parents=True,exist_ok=True)
        self.db=self.root/'tasks.sqlite3';self.lock=threading.RLock();self.export_lock=threading.Lock()
        with self.connect() as db:
            db.executescript('CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, payload TEXT NOT NULL); CREATE TABLE IF NOT EXISTS versions(id TEXT PRIMARY KEY, task_id TEXT NOT NULL, payload TEXT NOT NULL);')
        for t in self.list_full():
            if t.get('busy'):
                t.update(busy=False,status='运行中断',error='上次运行中断，已保存材料与旧版本，可重新运行。')
                self.save(t)
        self.consolidate_duplicates()

    def connect(self): return sqlite3.connect(self.db,timeout=15)
    def save(self,t):
        t['updated_at']=now()
        with self.connect() as db: db.execute('INSERT OR REPLACE INTO tasks VALUES (?,?)',(t['id'],json.dumps(t,ensure_ascii=False)))
    def get(self,tid,include_deleted=False):
        with self.connect() as db: r=db.execute('SELECT payload FROM tasks WHERE id=?',(tid,)).fetchone()
        if not r: raise AppError('任务不存在。',404)
        t=json.loads(r[0])
        if t.get('deleted_at') and not include_deleted: raise AppError('任务已删除，可在“已删除任务”中恢复。',404)
        return t
    def list_full(self):
        with self.connect() as db: return [json.loads(r[0]) for r in db.execute('SELECT payload FROM tasks')]
    @staticmethod
    def name_key(title,company):
        normalize=lambda value: re.sub(r'\s+','',unicodedata.normalize('NFKC',str(value))).casefold()
        return normalize(title),normalize(company)
    @staticmethod
    def case_id(t):
        return t.get('reference_case') or (t.get('generation_origin') or {}).get('case')
    @staticmethod
    def preferred(tasks):
        return max(tasks,key=lambda t:(bool(t.get('brief')) and t.get('reviewed_revision')==t['revision'],t['revision'],t['updated_at']))
    def find_existing(self,title,company,cid=None,exclude=None):
        key=self.name_key(title,company)
        matches=[t for t in self.list_full() if not t.get('deleted_at') and t['id']!=exclude and
                 ((cid and self.case_id(t)==cid) or self.name_key(t['title'],t['company'])==key)]
        return self.preferred(matches) if matches else None
    def consolidate_duplicates(self):
        """Archive duplicate records intact; never overwrite the retained draft."""
        with self.lock:
            tasks=[t for t in self.list_full() if not t.get('deleted_at')]
            if any(t.get('busy') for t in tasks): raise AppError('任务运行中，暂不整理重复记录。',409)
            groups=[]
            for t in tasks:
                keys={('name',self.name_key(t['title'],t['company']))}
                if self.case_id(t):keys.add(('case',self.case_id(t)))
                overlaps=[g for g in groups if g[0]&keys]
                members=[t]
                for g in overlaps:keys|=g[0];members+=g[1];groups.remove(g)
                groups.append((keys,members))
            archived=[]
            with self.connect() as db:
                for _,members in groups:
                    if len(members)<2:continue
                    keep=self.preferred(members)
                    for t in members:
                        if t['id']==keep['id']:continue
                        t.update(deleted_at=now(),duplicate_of=keep['id'])
                        self.event(t,'重复任务已归档，保留材料与版本；工作区使用 '+keep['title'])
                        db.execute('UPDATE tasks SET payload=? WHERE id=?',(json.dumps(t,ensure_ascii=False),t['id']))
                        archived.append({'id':t['id'],'kept':keep['id']})
            return archived
    def list(self,deleted=False):
        return sorted([{'id':t['id'],'title':t['title'],'company':t['company'],'status':t['status'],'revision':t['revision'],'updated_at':t['updated_at'],'source_count':len(t['sources']),'reference_case':t.get('reference_case'),'case_id':self.case_id(t),'duplicate_of':t.get('duplicate_of'),'has_brief':bool(t['brief'])} for t in self.list_full() if bool(t.get('deleted_at'))==deleted],key=lambda t:t['updated_at'],reverse=True)
    def delete(self,tid,revision=None):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            t['deleted_at']=now();self.event(t,'删除任务，材料与版本保留，可恢复');self.save(t)
            return {'id':tid,'deleted':True}
    def undelete(self,tid):
        with self.lock:
            t=self.get(tid,include_deleted=True)
            existing=self.find_existing(t['title'],t['company'],self.case_id(t),exclude=tid)
            if existing:return {**self.view(existing['id']),'reused_existing':True}
            if t.get('deleted_at'):
                t.pop('deleted_at');t.pop('duplicate_of',None);self.event(t,'恢复已删除任务');self.save(t)
            return self.view(tid)
    def event(self,t,text): t['events'].append({'at':now(),'text':text})
    def new(self,title,company=''):
        if not str(title).strip() or len(str(title))>120: raise AppError('请填写120字以内的任务名称。')
        with self.lock:
            existing=self.find_existing(title,str(company)[:80])
            if existing:return {**existing,'reused_existing':True}
            return self._new(title,company)
    def _new(self,title,company=''):
        t={'id':uid(),'title':str(title).strip(),'company':str(company)[:80],'status':'待上传','revision':0,'sources':[], 'notes':[], 'brief':None,'answers':{},'questions':[],'events':[],'busy':False,'error':'','created_at':now(),'updated_at':now(),'scope_authorization':'','reviewed_revision':None,'exports':[]}
        self.event(t,'创建任务');self.save(t);return t
    def ensure_idle(self,t,revision=None):
        if t['busy']: raise AppError('任务正在运行，请完成后再修改。',409)
        if revision is not None and revision!=t['revision']: raise AppError('任务已有更新，请刷新后再保存。',409)
    def changed(self,t):
        t['revision']+=1;t['reviewed_revision']=None;t['error']=''
    def snapshot(self,t,label):
        v={'id':uid(),'task_id':t['id'],'created_at':now(),'label':label,'revision':t['revision'],'brief':clone(t['brief']),'answers':clone(t['answers']),'scope_authorization':t['scope_authorization']}
        with self.connect() as db: db.execute('INSERT INTO versions VALUES (?,?,?)',(v['id'],t['id'],json.dumps(v,ensure_ascii=False)))
        return v
    def versions(self,tid):
        self.get(tid)
        with self.connect() as db: vs=[json.loads(r[0]) for r in db.execute('SELECT payload FROM versions WHERE task_id=?',(tid,))]
        vs.sort(key=lambda v:v['created_at'])
        for i,v in enumerate(vs,1):v['number']=i
        return list(reversed(vs))
    def view(self,tid):
        t=self.get(tid,include_deleted=True)
        if t.get('deleted_at'):
            existing=self.find_existing(t['title'],t['company'],self.case_id(t)) if t.get('duplicate_of') else None
            if existing:return {**self.view(existing['id']),'redirected_from':tid}
            raise AppError('任务已删除，可在“已删除任务”中恢复。',404)
        t['validation']=checked(t['brief']) if t['brief'] else []
        d=clone(t['brief']) if t['brief'] else {};d['confirmation_answers']=t['answers']
        t['confirmations']=confirmation_rows(d) if t['brief'] else []
        t['versions']=[{k:v[k] for k in ['id','created_at','label','revision','number']} for v in self.versions(tid)]
        t['rubric']=rubric_data()
        return t
    def add_file(self,tid,name,encoded):
        name=Path(name.replace('\\','/')).name
        if len(name)>180: raise AppError('文件名过长。')
        try: data=base64.b64decode(encoded,validate=True)
        except Exception: raise AppError('文件编码无效。')
        if not data or len(data)>12_000_000: raise AppError('单个文件需小于12MB且不能为空。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t)
            if len(t['sources'])>=20: raise AppError('每个任务最多20个文件。')
            digest=hashlib.sha256(data).hexdigest()
            if any(s.get('sha256')==digest for s in t['sources']): raise AppError('此文件已经上传，无需重复添加。')
            fid=uid();dest=self.root/'files'/tid/(fid+Path(name).suffix.lower());dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(data)
            try: text=extract_file(dest,name)
            except Exception as e:
                dest.unlink(missing_ok=True)
                if isinstance(e,AppError): raise
                raise AppError('文件无法解析，请检查文件是否完整、未加密且格式正确。') from None
            t['sources'].append({'id':fid,'name':name,'text':text,'path':str(dest),'sha256':digest,'role':'source','at':now()})
            self.changed(t);t['pending_update']=True;t['status']='待更新' if t['brief'] else '待拆解';self.event(t,'已读取材料：'+name);self.save(t)
        return self.view(tid)
    def add_note(self,tid,text,revision=None,scope=False,user_instruction=False):
        if not isinstance(text,str) or not text.strip() or len(text)>30000: raise AppError('请填写3万字以内的补充说明。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            t['notes'].append({'id':'U'+str(len(t['notes'])+1),'text':text.strip(),'at':now(),'scope_authorization':bool(scope),'user_instruction':bool(user_instruction)})
            if scope or user_instruction:
                t['scope_authorization']='仅允许执行以下用户说明中明确提出的范围调整，未提及的原要求继续保留：\n'+'\n'.join(n['text'] for n in t['notes'] if n.get('scope_authorization') or n.get('user_instruction'))
            self.changed(t);t['pending_update']=True;t['status']='待拆解' if not t['brief'] else '待更新';self.event(t,'已保存范围调整授权' if scope else '已保存补充说明');self.save(t)
        return self.view(tid)
    def set_brief(self,tid,brief,revision,label='手工修改'):
        errors=checked(brief)
        if errors: raise AppError('结构检查未通过：\n'+'\n'.join(errors))
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if brief['task_fit'].get('scope_change_applied') and not t['scope_authorization']: raise AppError('请先通过补充说明记录明确的范围调整授权。')
            brief=clone(brief);preserve_confirmation_controls(t.get('brief'),brief)
            if rubric_changed(t.get('brief'),brief):
                brief=clone(brief);brief['rubric_approved']=False;t['answers'].pop('C-RUBRIC',None)
            t['brief']=clone(brief);t['brief']['confirmation_answers']=t['answers']
            self.changed(t);t['pending_update']=False;t['status']='待审核';self.snapshot(t,label);self.event(t,label);self.save(t)
        return self.view(tid)
    def add_confirmation(self,tid,topic,question,revision):
        if not isinstance(topic,str) or not topic.strip() or len(topic)>80:raise AppError('请填写80字以内的确认事项标题。')
        if not isinstance(question,str) or not question.strip() or len(question)>4000:raise AppError('请填写4000字以内的待确认内容。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if not t['brief']:raise AppError('请先生成任务书。')
            t['brief'].setdefault('custom_confirmations',[]).append({'id':'C-USER-'+uid(),'section':'08','topic':topic.strip(),'question':question.strip(),'answer':''})
            self.changed(t);t['status']='待更新' if t.get('pending_update') else '待审核'
            self.snapshot(t,'添加企业确认事项');self.event(t,'添加企业确认事项：'+topic.strip());self.save(t)
        return self.view(tid)
    def delete_confirmation(self,tid,cid,revision):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            row=next((r for r in confirmation_rows(t['brief'] or {}) if r['id']==cid),None)
            if not row:raise AppError('确认事项不存在。',404)
            self.snapshot(t,'删除确认事项前的记录')
            custom=t['brief'].get('custom_confirmations',[])
            if any(c['id']==cid for c in custom):t['brief']['custom_confirmations']=[c for c in custom if c['id']!=cid]
            else:t['brief'].setdefault('dismissed_confirmations',[]).append(cid)
            t['answers'].pop(cid,None)
            if cid=='C-FORMAT':t['brief']['selected_output']=''
            if cid=='C-RUBRIC':t['brief']['rubric_approved']=False
            t['brief']['confirmation_answers']=t['answers']
            self.changed(t);t['status']='待更新' if t.get('pending_update') else '待审核'
            self.snapshot(t,'删除企业确认事项：'+row['topic']);self.event(t,'删除企业确认事项：'+row['topic']);self.save(t)
        return self.view(tid)
    def answer(self,tid,cid,text,by,revision,selection=None,approved=None):
        if not isinstance(text,str) or len(text)>4000: raise AppError('确认答案需在4000字以内。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            valid={r['id'] for r in confirmation_rows(t['brief'] or {})}
            if cid not in valid: raise AppError('确认事项不存在。')
            if cid=='C-FORMAT' and selection is not None:
                if selection and selection not in {o['id'] for o in t['brief']['output_options']}: raise AppError('输出格式选项不存在。')
                t['brief']['selected_output']=selection
            if cid=='C-RUBRIC' and approved is not None:
                if type(approved) is not bool: raise AppError('评分确认状态格式错误。')
                t['brief']['rubric_approved']=approved
            t['answers'][cid]={'text':text.strip(),'by':str(by or '内部产品负责人')[:80],'at':now()}
            t['brief']['confirmation_answers']=t['answers'];self.changed(t);t['status']='待更新' if t.get('pending_update') else '待审核'
            self.snapshot(t,'更新确认事项');self.event(t,'已记录确认答案；如影响业务内容，请补充修改要求后重新拆解。');self.save(t)
        return self.view(tid)
    def review(self,tid,revision):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if not t['brief'] or checked(t['brief']): raise AppError('需要有效的任务书草稿才能审核。')
            if t.get('pending_update') or t['status']=='待更新': raise AppError('存在尚未应用的材料或补充说明，请先更新任务书。')
            t['reviewed_revision']=t['revision'];t['status']='内部已审核';self.event(t,'已确认当前任务书，可导出 Word 和 Excel。');self.save(t)
        return self.view(tid)
    def restore(self,tid,vid,revision):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            v=next((v for v in self.versions(tid) if v['id']==vid),None)
            if not v: raise AppError('版本不存在。',404)
            t['brief']=clone(v['brief']);t['answers']=clone(v['answers']);t['scope_authorization']=v.get('scope_authorization','')
            self.changed(t);t['status']='待审核';self.snapshot(t,'恢复历史版本');self.event(t,'已恢复任务书版本；原始资料和补充记录仍保留。');self.save(t)
        return self.view(tid)
    def start(self,tid,revision,update_text=None):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if update_text is not None:
                if not isinstance(update_text,str):raise AppError('补充说明须为文字。')
                if update_text.strip() and not (t.get('pending_update') and t['notes'] and t['notes'][-1]['text']==update_text.strip()):
                    self.add_note(tid,update_text,revision,user_instruction=True);t=self.get(tid)
            if not config_status()['ready']: raise AppError('尚未接入模型。请技术同事配置接口地址、模型名称和API密钥后再运行；材料已保存。',409)
            if not t['sources'] and not t['notes']: raise AppError('请先上传原需求或填写需求说明。')
            if sum(len(s['text']) for s in t['sources'])+sum(len(n['text']) for n in t['notes'])>120000: raise AppError('本任务材料合计超过12万字，请拆分任务后运行。')
            t.update(busy=True,status='正在更新' if t.get('brief') else '正在拆解',error='');self.event(t,'正在根据补充信息更新任务书' if t.get('brief') else '正在理解材料和检查缺口');self.save(t)
        threading.Thread(target=self.run,args=(tid,),daemon=True).start();return self.view(tid)
    def run(self,tid):
        calls=[];started_at=now()
        def record(usage,outcome='returned'):
            calls.append({'at':now(),'usage':usage,'outcome':outcome})
        def stamp(t):
            t['usage']={k:sum(c['usage'][k] for c in calls) if calls and all(isinstance(c['usage'].get(k),(int,float)) for c in calls) else None for k in ['prompt_tokens','completion_tokens','total_tokens']}
            t['last_run']={'started_at':started_at,'finished_at':now(),'model':config_status()['model'],'calls':calls}
            t.setdefault('runs',[]).append(clone(t['last_run']))
        try:
            original=self.get(tid);messages=messages_for(original)
            result=None;usage={}
            for attempt in range(2):
                if attempt:
                    with self.lock:
                        t=self.get(tid);self.event(t,'草稿需要修正结构，正在保留业务要求并修复');self.save(t)
                try: result,usage=call_model(messages,response_kind='repair' if attempt else ('revision' if original.get('brief') else 'draft'))
                except ModelError as e:
                    record(e.usage,'error')
                    if attempt==0 and e.response_text is not None:
                        messages.extend([{'role':'assistant','content':e.response_text}, {'role':'user','content':'返回格式未满足工作台协议：'+str(e)+' 请保留原需求，返回一个JSON对象，顶层action只能是draft或clarify；draft时完整任务书放在brief中。'}])
                        continue
                    raise
                record(usage)
                result=clone(result)
                if result['action']=='clarify':
                    qs=result.get('questions')
                    if not isinstance(qs,list) or not 1<=len(qs)<=3 or any(not isinstance(q,dict) or not isinstance(q.get('question'),str) or not q['question'].strip() or not isinstance(q.get('id'),str) for q in qs): raise ModelError('模型返回的问题格式无效，请重试。')
                    break
                d=result.get('brief');errors=checked(d)
                if isinstance(d,dict) and 'rules' not in d: errors.append('缺少rules数组：原文业务规则、禁止项和范围边界需在交付正文可见，不能只存于runtime或task_fit；无业务规则时才可为空数组')
                if isinstance(d,dict) and d.get('task_fit',{}).get('scope_change_applied') and not original['scope_authorization']: errors.append('不能擅自应用范围调整；先提问取得用户授权。')
                if isinstance(d,dict) and d.get('rubric_approved') and not(original.get('brief') or {}).get('rubric_approved'): errors.append('不能自动把评分细则标为企业已批准。')
                if isinstance(d,dict) and d.get('selected_output','')!=(original.get('brief') or {}).get('selected_output',''): errors.append('输出格式选择必须由用户在确认事项中记录，不得自行选定或更改。')
                if isinstance(d,dict) and (original.get('brief') or {}).get('rubric_approved') and any(d.get(k)!=(original.get('brief') or {}).get(k) for k in ['outputs','kr2_2','kr2_3']):
                    d['rubric_approved']=False
                calls[-1]['validation_errors']=clone(errors)
                if not errors: break
                if attempt==1: raise ModelError('两次结构检查仍未通过：'+'；'.join(errors[:8]))
                messages.extend([{'role':'assistant','content':json.dumps(result,ensure_ascii=False)}, {'role':'user','content':'只修复这些结构问题并保留业务要求：'+json.dumps(errors,ensure_ascii=False)}])
            quality=[]
            if result['action']=='draft':
                result,compact_issues=self.compact_draft(original,result,record)
                result,quality=self.audit_and_repair(original,result,record)
                quality=list(quality)+compact_issues
            with self.lock:
                t=self.get(tid);t['busy']=False;stamp(t);t['quality_review']=list(quality);self.changed(t)
                if result['action']=='clarify': t['questions']=result['questions'];t['status']='待补充'
                else:
                    if rubric_changed(original.get('brief'),result['brief']):
                        result['brief']['rubric_approved']=False
                        t['answers'].pop('C-RUBRIC',None)
                    preserve_confirmation_controls(original.get('brief'),result['brief'])
                    t['brief']=result['brief'];t['brief']['confirmation_answers']=t['answers'];t['pending_update']=False;t['questions']=[];t['status']='待审核';self.snapshot(t,'AI生成' if not original['brief'] else 'AI修改')
                self.event(t,str(result.get('summary','本次处理已完成'))[:1000]);self.save(t)
        except Exception as e:
            with self.lock:
                t=self.get(tid);stamp(t);t.update(busy=False,status='运行失败',error=str(e) if isinstance(e,(ModelError,AppError)) else '处理失败，已保留资料与旧版本。请检查配置后重试。');self.event(t,t['error']);self.save(t)
    def compact_draft(self,original,result,record):
        d=result['brief'];old=original.get('brief')
        fields=('taxonomy','scenario','purpose','inputs','rules','outputs','objective','kr2_2')
        modes={k:'full' for k in fields if not old or d.get(k)!=old.get(k)}
        if not modes:return result,[]
        if 'inputs' in modes or 'outputs' in modes:
            modes.setdefault('kr2_2','refs');modes['confirmations']='refs'
        with self.lock:
            t=self.get(original['id']);self.event(t,'正在合并重复输入、交付与评分，保留必要业务要求');self.save(t)
        try:
            response,usage=compact_brief(original,d,modes);record(usage,'content_compact')
            changes=response.get('brief')
            if response.get('action')!='draft' or not isinstance(changes,dict) or set(changes)-set(modes):
                raise ValueError('精简步骤返回了未授权字段')
            for key,value in changes.items():
                if modes[key]=='refs':
                    without_refs=lambda rows:[{k:v for k,v in x.items() if k!='refs'} for x in rows]
                    if without_refs(value)!=without_refs(d.get(key,[])):raise ValueError('精简步骤改动了无关正文')
            candidate=clone(d);candidate.update(changes)
            errors=checked(candidate)
            if errors:raise ValueError('；'.join(errors[:3]))
            result=clone(result);result['brief']=candidate
            return result,[]
        except (ModelError,ValueError,TypeError,AttributeError) as e:
            if isinstance(e,ModelError):record(e.usage,'error')
            return result,[{'problem':'任务书精简未完成','source':'必要性整理步骤','fix':'保留有效草稿，请检查重复内容。'+str(e)}]

    def audit_and_repair(self,original,result,record):
        quality=[]
        for attempt in range(4):
            with self.lock:
                t=self.get(original['id']);self.event(t,'正在对照原需求和后续补充核查内容'+('（修正后复核）' if attempt else ''));self.save(t)
            feedback=None
            for connection_try in range(2):
                try:
                    quality,usage=audit_brief(original,result['brief'],feedback=feedback,prior_issues=list(quality));record(usage,'content_audit')
                    with self.lock:
                        t=self.get(original['id']);t.setdefault('content_audits',[]).append({'at':now(),'brief_sha256':hashlib.sha256(json.dumps(result['brief'],ensure_ascii=False,sort_keys=True).encode()).hexdigest(),'report':getattr(quality,'report',{'issues':list(quality)}),'passed':not quality});self.save(t)
                    break
                except ModelError as e:
                    record(e.usage,'error')
                    if connection_try==0 and (e.transient or e.response_text is not None):
                        if e.response_text is not None:feedback={'response':e.response_text,'error':str(e)}
                        with self.lock:
                            t=self.get(original['id']);self.event(t,'内容复核连接中断，保留草稿并重试一次' if e.transient else '复核报告的来源引用需要修正，正在重新核对');self.save(t)
                        continue
                    return result,[{'problem':'本次内容复核未完成','source':'模型复核步骤','fix':str(e)+' 请对照原始材料人工审核。'}]
            if not quality or attempt==3: break
            repair_messages=messages_for(original)+[
                {'role':'assistant','content':json.dumps(result,ensure_ascii=False)},
                {'role':'user','content':'独立复核发现以下与原材料或补充意见不一致的问题。核对依据，仅修正这些问题及直接关联项，其他内容不改。仍返回完整draft JSON；不能自行批准或选择格式。\n'+json.dumps(quality,ensure_ascii=False)}]
            try:
                repaired,usage=call_model(repair_messages,response_kind='revision');record(usage,'content_repair')
            except ModelError as e:
                record(e.usage,'error');return result,quality+[{'problem':'自动修正未完成','source':'模型修正步骤','fix':str(e)}]
            d=repaired.get('brief');errors=checked(d)
            if isinstance(d,dict) and 'rules' not in d:errors.append('修复后缺少rules')
            prior=original.get('brief') or {}
            if repaired.get('action')!='draft' or errors or (d.get('task_fit',{}).get('scope_change_applied') and not original['scope_authorization']) or d.get('selected_output','')!=prior.get('selected_output','') or (d.get('rubric_approved') and not prior.get('rubric_approved')):
                return result,quality+[{'problem':'自动修正未通过结构或授权检查','source':'系统检查','fix':'已保留修正前有效草稿，请人工修改。'+ '；'.join(errors[:3])}]
            if prior.get('rubric_approved') and any(d.get(k)!=prior.get(k) for k in ['outputs','kr2_2','kr2_3']):d['rubric_approved']=False
            result=repaired
        return result,quality
    def build_word(self,d,path):
        word_builder.build(delivery_data(d),path)
        self.append_rules(path,d);self.finish_layout(path,d)
        from docx import Document
        Document(path)
    def version_document(self,tid,vid):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t)
            v=next((v for v in self.versions(tid) if v['id']==vid),None)
            if not v or not v.get('brief'):raise AppError('历史版本不存在。',404)
            d=clone(v['brief']);d['confirmation_answers']=clone(v['answers'])
            if checked(d):raise AppError('该历史版本的数据不完整，暂时无法生成 Word 预览。')
            folder=self.root/'previews'/tid/vid/uid();folder.mkdir(parents=True)
            path=folder/('企业任务书_V'+str(v['number'])+'.docx')
            with self.export_lock:self.build_word(d,path)
            return path,path.name
    def export(self,tid,revision,kind='both'):
        if kind not in ('both','docx','xlsx'):raise AppError('导出格式无效。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if t['reviewed_revision']!=t['revision']: raise AppError('请先确认当前任务书版本，再生成交付文件。')
            d=clone(t['brief']);errors=checked(d)
            if errors: raise AppError('\n'.join(errors))
            d['confirmation_answers']=t['answers'];eid=uid();out=self.root/'exports'/tid/eid;out.mkdir(parents=True)
            stem=re.sub(r'[\\/:*?"<>|\x00-\x1f]','_',d['title'])[:65]
            files=[]
            with self.export_lock:
                if kind in ('both','docx'):
                    path=out/(stem+'_任务书.docx');self.build_word(d,path);files.append(path)
                if kind in ('both','xlsx'):
                    path=out/(stem+'_输入样例表.xlsx');excel_builder.build(delivery_data(d),path);self.finish_sheet(path,d)
                    from openpyxl import load_workbook
                    w=load_workbook(path,read_only=True);w.close();files.append(path)
            versions=self.versions(tid)
            number=next((v['number'] for v in versions if v['revision']==t['revision']),None)
            item={'id':eid,'revision':t['revision'],'version_number':number,'created_at':now(),'files':[{'name':p.name,'kind':p.suffix[1:],'path':str(p)} for p in files],'visual_review':'pending'}
            t['exports'].append(item);self.event(t,{'docx':'已生成企业任务书 Word','xlsx':'已生成企业输入样例表 Excel','both':'已生成 Word 和 Excel'}[kind]);self.save(t)
        return self.view(tid)
    def finish_sheet(self,path,d):
        from openpyxl import load_workbook
        wb=load_workbook(path);ws=wb.active
        count=d.get('sample',{}).get('count',3)
        ws['A2']=sample_request(d)+' 选填项可留空；不填写密码、Cookie、验证码或密钥。'
        ws.row_dimensions[2].height=48
        for row,x in enumerate(d['inputs'],5):
            fallback=x.get('student_fallback')
            if fallback:
                sentence=fallback if '企业' in fallback or '学生' in fallback or '参与者' in fallback else '企业未提供时，由学生'+fallback
                detail=x['requirement'].rstrip('。；;')+'；'+sentence
                if x.get('dependency_note'): detail+='（'+x['dependency_note']+'）'
                ws.cell(row,3,excel_builder.literal(detail))
        wb.save(path);wb.close()
    def append_rules(self,path,d):
        if not d.get('rules'): return
        from docx import Document
        doc=Document(path)
        anchor=next((p for p in doc.paragraphs if p.text=='04 希望得到什么？'),None)
        # Pinned heading contains two spaces; match the semantic heading.
        if anchor is None: anchor=next(p for p in doc.paragraphs if p.text.startswith('04'))
        writer=word_builder.Writer(d['title']);writer.doc=doc
        p=writer.para('数据处理规则',10.5,True);anchor._p.addprevious(p._p)
        table=writer.table(['规则','具体要求'],[[r['name'],r['detail']] for r in d['rules']],[2,8])
        anchor._p.addprevious(table._tbl)
        word_builder.clear_nonheading_markers(doc);doc.save(path)

    def finish_layout(self,path,d=None):
        from docx import Document
        doc=Document(path)
        for x in (d or {}).get('inputs',[]):
            fallback=x.get('student_fallback','')
            if fallback and any(w in fallback for w in ['企业','学生','参与者']):
                old='；企业未提供时，由学生'+fallback
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    if old in run.text:run.text=run.text.replace(old,'；'+fallback)
        # Run last: non-heading keepNext/keepLines settings appear as black
        # paragraph-format markers in Word. Preserve heading styles, repeated
        # table headers and non-splitting rows without reintroducing these marks.
        word_builder.clear_nonheading_markers(doc)
        # Start the confirmation sheet on a new page using its heading only;
        # the introduction and table cells remain ordinary body paragraphs.
        for p in doc.paragraphs:
            if p.style.name == 'Heading 1' and p.text.startswith('08'):
                p.paragraph_format.page_break_before=True
        doc.save(path)

    def cases(self):
        p=ROOT/'research/cases.json'
        return json.loads(p.read_text()) if p.exists() else []
    def case_file(self,cid,index):
        c=next((c for c in self.cases() if c['id']==cid),None)
        if not c or not str(index).isdigit() or int(index)>=len(c['files']): raise AppError('文件不存在。',404)
        p=Path(c['files'][int(index)]['path']).resolve()
        if not p.is_relative_to((ROOT/'research/materials').resolve()) or p.suffix.lower() not in {'.md','.docx','.xlsx','.pdf','.txt','.csv'} or not p.is_file(): raise AppError('原文件不可用。',404)
        return p,p.name
    def case_detail(self,cid):
        c=next((c for c in self.cases() if c['id']==cid),None)
        if not c: raise AppError('案例不存在。',404)
        c=clone(c)
        for f in c['files']:
            p=Path(f.get('extracted',''))
            f['text']=p.read_text() if p.is_file() else ''
        return c
    def import_case(self,cid):
        with self.lock:
            c=self.case_detail(cid)
            existing=self.find_existing(c['title'],c.get('company',''),cid)
            if existing:return {**self.view(existing['id']),'reused_existing':True}
            deleted=[t for t in self.list_full() if t.get('deleted_at') and self.case_id(t)==cid]
            if deleted:
                t=self.undelete(self.preferred(deleted)['id'])
                return {**t,'restored_existing':True}
            return self._import_case(c)
    def _import_case(self,c):
        cid=c['id'];t=self._new(c['title'],c.get('company',''))
        t['reference_case']=cid;t['reference_note']='来自用户提供的历史材料；不是本系统新生成的结果。'
        for f in c['files']:
            p=Path(f['path'])
            if not p.is_file(): continue
            fid=uid();dest=self.root/'files'/t['id']/(fid+p.suffix);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_bytes(p.read_bytes())
            t['sources'].append({'id':fid,'name':p.name,'text':f['text'],'path':str(dest),'role':'source' if f['role']=='source' else 'reference','sha256':f['sha256'],'at':now()})
        model_file=ROOT/'research/reference-data'/f'{cid}.json'
        if model_file.exists():
            d=json.loads(model_file.read_text())
            if not checked(d):
                t['brief']=d;t['scope_authorization']=d.get('task_fit',{}).get('conversion_note','') if d.get('task_fit',{}).get('scope_change_authorized') else ''
                t['status']='待审核';t['revision']=1;self.snapshot(t,'导入历史参考稿')
        if not t['brief']: t['status']='待拆解'
        notes_file=ROOT/'research/case-updates.json'
        if notes_file.exists():
            for item in json.loads(notes_file.read_text()).get(cid,[]): t['notes'].append(item)
        authorized=[n['text'] for n in t['notes'] if n.get('scope_authorization')]
        if authorized: t['scope_authorization']='\n'.join(authorized)
        self.event(t,'导入历史案例，参考成稿与原需求已区分');self.save(t);return self.view(t['id'])

    def file(self,tid,fid):
        t=self.get(tid)
        f=next((f for f in t['sources'] if f['id']==fid),None)
        if not f:
            for e in t['exports']:
                for item in e['files']:
                    if fid==e['id']+'-'+item['kind']: f=item
        if not f: raise AppError('文件不存在。',404)
        path=Path(f['path']).resolve()
        if not path.is_relative_to(self.root) or not path.is_file(): raise AppError('文件不可用。',404)
        return path,f['name']
