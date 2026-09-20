"""Resumable, bounded task-book agent. Existing workspace remains independently usable."""
import hashlib
import json
import threading
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import core
from core import Store, AppError, clone, now, uid, checked, rubric_changed, preserve_confirmation_controls
from model import call_model, messages_for, audit_brief, compact_brief, config_status, ModelError
from confirmation_actions import CONTRACT_VERSION, decisions, apply_decisions

FIELDS=('title','taxonomy','scenario','purpose','inputs','rules','outputs','objective','kr2_2','kr2_3','runtime','task_fit','scope_assessment','sample','output_options','confirmations','sources','reviewer')
LABELS={'title':'任务名称','taxonomy':'行业与岗位','scenario':'01 场景','purpose':'02 学生交付物','inputs':'03 输入','rules':'处理规则','outputs':'04 输出与 KR1','objective':'05 目标','kr2_2':'KR2 业务质量','kr2_3':'KR3 使用表现','runtime':'使用边界','task_fit':'任务范围','scope_assessment':'范围说明','sample':'输入样例','output_options':'输出格式选项','confirmations':'待确认事项','sources':'来源依据'}
LABELS['reviewer']='评分人'
STAGES={'plan':'理解材料与修改意见','generate':'生成任务书','compact':'整理必要内容','audit':'核对来源与要求','repair':'修正检查发现的问题','commit':'保存可审阅版本','ready':'等待你审阅','questions':'等待必要回答','done':'交付文件已生成'}
MAX_CALLS=8
TOKEN_STOP=200000

class AgentPause(Exception):pass

def basis(t):
    value={k:clone(t.get(k)) for k in ('revision','sources','notes','brief','answers','scope_authorization')}
    # A cloud cache may move on every restart. File content and identity remain
    # part of the basis; temporary filesystem locations do not.
    for source in value.get('sources') or []:
        source.pop('path',None);source.pop('cloud_saved',None)
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True).encode()).hexdigest()

def input_task(t):
    return {k:clone(t.get(k)) for k in ('id','title','company','sources','notes','brief','answers','scope_authorization','pending_update')}

def questions(value):
    if not isinstance(value,list) or not 1<=len(value)<=3:raise ModelError('需要补充的问题应为1至3个。')
    if any(not isinstance(q,dict) or not isinstance(q.get('id'),str) or not isinstance(q.get('question'),str) or not q['question'].strip() for q in value):raise ModelError('问题格式不完整，已保留进度。')
    return value

class AgentStore(Store):
    def __init__(self,data=None):
        super().__init__(data)
        for t in self.list_full():
            a=t.get('agent')
            if a and a.get('status')=='running':
                a.update(status='paused',pause_requested=False,message='运行中断，已保存完成步骤；点击继续处理。')
                t.update(busy=False,status='已暂停');self.save(t)

    def view(self,tid):
        t=super().view(tid)
        if t.get('agent'):
            t['agent']['stale']=t['agent'].get('basis')!=basis(t)
        return t

    def start_agent(self,tid,revision,text='',resume=False,reply_to=None):
        if not isinstance(text,str) or len(text)>30000:raise AppError('请填写3万字以内的补充信息。')
        if type(revision) is not int or type(resume) is not bool:raise AppError('处理请求格式错误，请刷新后重试。')
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision);text=text.strip()
            old=t.get('agent',{});preexisting_pending=bool(t.get('pending_update'))
            if reply_to is not None and reply_to!=old.get('id'):raise AppError('问题已有更新，请刷新后回答。',409)
            if text:
                qs=t.get('questions',[])
                self.add_note(tid,text,revision,user_instruction=True);t=self.get(tid)
                if qs:
                    # Assistant questions give context, never become user scope authorization.
                    t['notes'][-1]['answering_questions']=clone(qs);self.save(t)
            if not t['sources'] and not t['notes']:raise AppError('先上传材料或写下需求。')
            if sum(len(s['text']) for s in t['sources'])+sum(len(n['text']) for n in t['notes'])>120000:raise AppError('材料合计超过12万字，请拆分任务。')
            if resume and not text and old.get('basis')==basis(t) and old.get('status') in ('paused','blocked'):
                a=clone(old);a.update(status='running',pause_requested=False,message='继续处理已保存的进度',session_calls=0,session_tokens=0)
                if a.get('contract_version')!=CONTRACT_VERSION:
                    # Reinterpret the saved instruction once after the confirmation contract changes.
                    a.update(phase='plan',plan=None,candidate=None,format_count=0,repair_count=0,format_feedback='',retry_counts={})
            else:
                a={'id':uid(),'status':'running','phase':'plan','basis':basis(t),'original':input_task(t),'direct_message':text,'preexisting_pending':preexisting_pending,
                   'candidate':None,'plan':None,'issues':[],'changes':[],'trace':[],'calls':[],'session_calls':0,'session_tokens':0,
                   'repair_count':0,'format_count':0,'retry_counts':{},'pause_requested':False,'started_at':now(),'message':'正在判断下一步'}
            a['contract_version']=CONTRACT_VERSION
            if not config_status()['ready']:
                a.update(status='blocked',message='尚未配置可用模型，资料和进度已保存。');t.update(agent=a,busy=False,status='已暂停',error=a['message']);self.save(t);return self.view(tid)
            t.update(agent=a,busy=True,status='正在更新' if t.get('brief') else '正在拆解',error='');self.event(t,'Agent 开始连续处理');self.save(t)
        threading.Thread(target=self.run_agent,args=(tid,a['id']),daemon=True).start()
        return self.view(tid)

    def pause_agent(self,tid,revision):
        with self.lock:
            t=self.get(tid)
            if revision!=t['revision']:raise AppError('任务已有更新，请刷新。',409)
            if not t.get('agent') or not t['busy']:return self.view(tid)
            t['agent']['pause_requested']=True;t['agent']['message']='将在当前请求返回后暂停，已完成步骤会保留。';self.save(t)
        return self.view(tid)

    def answer_and_continue(self,tid,revision,body):
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            row=next((r for r in core.confirmation_rows(t.get('brief') or {}) if r['id']==body.get('id')),None)
            if not row:raise AppError('确认事项不存在。')
            result=self.answer(tid,body['id'],body.get('text',''),body.get('by'),revision,body.get('selection'),body.get('approved'))
            if body['id'] in ('C-FORMAT','C-RUBRIC'):
                # Controls are deterministic: no model is needed to copy a choice.
                t=self.get(tid)
                if t.get('agent',{}).get('status')=='done':t['agent'].update(status='ready',phase='ready',message='确认信息已更新，请重新确认并生成文件。')
                self.save(t);return self.view(tid)
            if not body.get('text','').strip():return result
            return self.start_agent(tid,result['revision'],'关于“'+row['topic']+'”的明确回答：'+body['text'])

    def _save_agent(self,tid,a):
        with self.lock:
            t=self.get(tid)
            if t.get('agent',{}).get('id')!=a['id']:raise AgentPause('任务已被新的处理替代。')
            a['pause_requested']=t['agent'].get('pause_requested',False)
            t['agent']=clone(a);self.save(t)

    def _guard(self,tid,a):
        t=self.get(tid)
        if t.get('agent',{}).get('id')!=a['id']:raise AgentPause('任务已有新处理。')
        if t['agent'].get('pause_requested'):raise AgentPause('已暂停，完成的步骤已保留。')
        if basis(t)!=a['basis']:raise AgentPause('任务内容已经变化，请按新内容重新处理。')

    def _call(self,tid,a,label,fn):
        self._guard(tid,a)
        if a['session_calls']>=MAX_CALLS or a['session_tokens']>=TOKEN_STOP:raise AgentPause('本轮达到调用或用量上限，进度已保存；可审阅后继续。')
        entry={'id':uid(),'step':label,'at':now(),'outcome':'pending','usage':{}}
        a['calls'].append(entry);a['session_calls']+=1;a['message']=STAGES.get(a['phase'],label);self._save_agent(tid,a)
        try:
            result,usage=fn();entry.update(outcome='returned',usage=usage)
            tokens=usage.get('total_tokens')
            if isinstance(tokens,(int,float)):a['session_tokens']+=max(0,tokens)
            return result
        except ModelError as e:
            entry.update(outcome='error',usage=e.usage)
            tokens=e.usage.get('total_tokens')
            if isinstance(tokens,(int,float)):a['session_tokens']+=max(0,tokens)
            raise
        finally:self._save_agent(tid,a)

    def _plan(self,tid,a):
        t=a['original'];payload={'task_title':t['title'],'source_materials':[{'name':s['name'],'text':s['text']} for s in t['sources'] if s.get('role')!='reference'],
            'user_updates':t['notes'],'current_brief':t['brief'],'confirmation_answers':t['answers'],'confirmation_items':core.confirmation_rows(t['brief']) if t['brief'] else [],'direct_message':a['direct_message'],'allowed_fields':FIELDS}
        r=self._call(tid,a,'判断下一步',lambda:call_model([{'role':'system','content':(core.ROOT/'prompts/next-step.md').read_text()},{'role':'user','content':json.dumps(payload,ensure_ascii=False)}],response_kind='plan'))
        if r.get('intent') not in ('draft','revise','record','clarify'):raise ModelError('下一步决策不符合约定。')
        fs=r.get('fields',[])
        if not isinstance(fs,list) or any(not isinstance(x,str) or x not in FIELDS for x in fs):raise ModelError('修改范围不符合约定。')
        if r['intent']=='revise' and not fs:raise ModelError('修改任务需要明确影响范围。')
        try:
            r['decisions']=decisions(t,r,a['direct_message'])
        except ModelError as e:
            # A malformed model answer must not force the user to repeat a clear instruction.
            a['invalid_plan']=clone(r)
            payload['invalid_plan']=r;payload['validation_feedback']=str(e)
            payload['correction_instruction']='仅修正确认动作：id 必须逐字使用 confirmation_items 的 id，每个 id 最多一次。批量确认只用 confirm_all；answers 仅列用户明确给出的具体答案或选择。所有 quote 必须是 direct_message 中逐字存在的原话。返回完整 plan。'
            r=self._call(tid,a,'修正确认动作',lambda:call_model([{'role':'system','content':(core.ROOT/'prompts/next-step.md').read_text()},{'role':'user','content':json.dumps(payload,ensure_ascii=False)}],response_kind='plan'))
            if r.get('intent') not in ('draft','revise','record','clarify'):raise ModelError('下一步决策不符合约定。')
            fs=r.get('fields',[])
            if not isinstance(fs,list) or any(not isinstance(x,str) or x not in FIELDS for x in fs):raise ModelError('修改范围不符合约定。')
            if r['intent']=='revise' and not fs:raise ModelError('修改任务需要明确影响范围。')
            r['decisions']=decisions(t,r,a['direct_message'])
        if r['intent']=='record' and (not t.get('brief') or a.get('preexisting_pending') or a['direct_message'] and not r['decisions']):raise ModelError('仍有未落实内容，不能只记录答案。')
        if any(item.get('option') for item in r['decisions']) and 'output_options' not in fs:fs.append('output_options')
        if r['intent']=='clarify':questions(r.get('questions'))
        if r['intent']=='draft' and t.get('brief'):r['intent']='revise';r['fields']=list(FIELDS)
        if r['intent']=='revise' and not t.get('brief'):r['intent']='draft'
        a['plan']=r
        if r['intent']=='clarify':a['phase']='questions'
        elif r['intent']=='record':a['candidate']=clone(t['brief']);a['phase']='commit'
        else:a['phase']='generate'

    def _model_task(self,a):
        t=clone(a['original'])
        if t.get('brief'):
            t['answers']=apply_decisions(t,t['brief'],a['plan'].get('decisions',[]))
        return t

    def _validate(self,a,d,reference=None):
        if not isinstance(d,dict):return checked(d)
        errors=[]
        old=a['original'].get('brief') or {}
        effective=self._model_task(a).get('brief') or {}
        if 'rules' not in d:errors.append('缺少处理规则数组')
        if d.get('selected_output','') not in (old.get('selected_output',''),effective.get('selected_output','')):errors.append('输出格式变更缺少用户明确选择')
        if d.get('rubric_approved') and not (old.get('rubric_approved') or effective.get('rubric_approved')):errors.append('不得代替企业确认评分')
        if errors:return errors
        # Controls and answers are written from grounded decisions, never free-form draft guesses.
        d['selected_output']=old.get('selected_output','');d['rubric_approved']=bool(old.get('rubric_approved'))
        apply_decisions(a['original'],d,a['plan'].get('decisions',[]))
        errors+=checked(d)
        if d.get('task_fit',{}).get('scope_change_applied') and not a['original'].get('scope_authorization'):errors.append('缺少范围调整授权')
        if old and a['plan']['intent']=='revise':
            allowed=set(a['plan']['fields'])|{'sources'}
            if allowed&{'inputs','outputs'}:allowed|={'kr2_2','confirmations'}
            for key in FIELDS:
                if key not in allowed and d.get(key)!=old.get(key):errors.append('不应改动未涉及的字段：'+key)
        return errors

    def _generate(self,tid,a):
        msgs=messages_for(self._model_task(a))
        if a['plan']['intent']=='revise':msgs.append({'role':'user','content':'本轮仅允许修改以下字段及其直接关联引用，其他字段逐字保留：'+json.dumps(a['plan']['fields'],ensure_ascii=False)})
        if a.get('format_feedback'):msgs.append({'role':'user','content':'上次输出未通过校验，请修复并重新输出完整任务书：'+a['format_feedback']})
        r=self._call(tid,a,'生成或局部修改',lambda:call_model(msgs,response_kind='revision' if a['original'].get('brief') else 'draft'))
        if r.get('action')=='clarify':
            a['plan']['questions']=questions(r.get('questions'));a['phase']='questions';return
        d=r.get('brief');errors=self._validate(a,d)
        if errors:
            a['format_count']+=1
            if a['format_count']>1:raise ModelError('任务书未通过校验：'+'；'.join(errors[:5]))
            a['format_feedback']='；'.join(errors);return
        a['candidate']=d;a['phase']='compact'

    def _compact(self,tid,a):
        d=a['candidate'];old=a['original'].get('brief');fields=('taxonomy','scenario','purpose','inputs','rules','outputs','objective','kr2_2')
        modes={k:'full' for k in fields if not old or d.get(k)!=old.get(k)}
        if not modes:a['phase']='audit';return
        if 'inputs' in modes or 'outputs' in modes:modes.setdefault('kr2_2','refs');modes['confirmations']='refs'
        r=self._call(tid,a,'精简重复内容',lambda:compact_brief(self._model_task(a),d,modes))
        patch=r.get('brief');err=[]
        if r.get('action')!='draft' or not isinstance(patch,dict) or set(patch)-set(modes):err=['精简返回了范围外字段']
        else:
            for k,v in patch.items():
                if modes[k]=='refs':
                    strip=lambda rows:[{key:value for key,value in row.items() if key!='refs'} for row in rows]
                    try:
                        if strip(v)!=strip(d.get(k,[])):err.append('精简改动了无关正文')
                    except (TypeError,AttributeError):err.append('精简引用格式错误')
            merged=clone(d);merged.update(patch);err+=self._validate(a,merged)
            if not err:a['candidate']=merged
        if err:a['issues'].append({'problem':'精简步骤未通过，保留原草稿','source':'结构检查','fix':'；'.join(err[:3])})
        a['phase']='audit'

    def _audit(self,tid,a):
        findings=self._call(tid,a,'核对业务要求',lambda:audit_brief(self._model_task(a),a['candidate'],prior_issues=a.get('findings',[])))
        a['findings']=list(findings)
        a['phase']='repair' if findings and a['repair_count']<2 else 'commit'

    def _repair(self,tid,a):
        msgs=messages_for(self._model_task(a))+[{'role':'assistant','content':json.dumps({'action':'draft','brief':a['candidate']},ensure_ascii=False)},
             {'role':'user','content':'只修正以下问题及直接关联内容，其他字段保持原样；返回完整draft，不代替用户批准：'+json.dumps(a['findings'],ensure_ascii=False)}]
        r=self._call(tid,a,'修正业务问题',lambda:call_model(msgs,response_kind='revision'));d=r.get('brief');err=self._validate(a,d)
        a['repair_count']+=1
        if r.get('action')!='draft' or err:raise ModelError('修正未通过校验，保留上一份草稿：'+'；'.join(err[:4]))
        a['candidate']=d;a['phase']='audit'

    def _finish(self,tid,a):
        with self.lock:
            self._guard(tid,a);t=self.get(tid)
            if a['phase']=='questions':
                t['questions']=a['plan']['questions'];t['status']='待补充';a.update(status='waiting',message='请回答必要问题，我会接着处理。')
                t['busy']=False;t['agent']=a;self.event(t,a['message']);self.save(t);return
            d=clone(a['candidate']);old=t.get('brief');preserve_confirmation_controls(old,d)
            errors=self._validate(a,d)
            if errors:raise ModelError('保存前校验未通过：'+'；'.join(errors[:5]))
            t['answers']=apply_decisions(a['original'],d,a['plan'].get('decisions',[]))
            a['changes']=[{'field':k,'label':LABELS[k],'before':clone((old or {}).get(k)),'after':clone(d.get(k))} for k in FIELDS if (old or {}).get(k)!=d.get(k)]
            d['confirmation_answers']=t['answers'];t['brief']=d;self.changed(t);t.update(pending_update=False,questions=[],quality_review=a.get('issues',[])+a.get('findings',[]),busy=False,status='待审核',error='')
            count=len(a['plan'].get('decisions',[]))
            message=f'已同步{count}项确认信息，请审阅任务书。' if count else '已完成处理，请审阅任务书。'
            a.update(status='ready',phase='ready',message=message,basis=basis(t),completed_at=now())
            t['agent']=a;self.snapshot(t,'Agent 生成' if not old else 'Agent 修改');self.event(t,a['message']);self.save(t)

    def run_agent(self,tid,run_id):
        a=clone(self.get(tid)['agent'])
        if a['id']!=run_id:return
        try:
            while a['phase'] not in ('ready','done'):
                self._guard(tid,a);phase=a['phase']
                if phase in ('commit','questions'):self._finish(tid,a);return
                a['message']=STAGES[phase];self._save_agent(tid,a)
                try:getattr(self,'_'+('plan' if phase=='plan' else phase))(tid,a)
                except ModelError as e:
                    key=phase+':'+str(a['repair_count']);retries=a['retry_counts'].get(key,0)
                    if e.transient and retries<1:
                        a['retry_counts'][key]=retries+1;a['trace'].append({'at':now(),'text':'连接中断，保留进度并重试一次'});self._save_agent(tid,a);continue
                    raise
                a['trace'].append({'at':now(),'text':STAGES[phase]+('已完成' if a['phase']!=phase else '需要修正格式，继续处理'),'next':a['phase']});self._save_agent(tid,a)
        except Exception as e:
            with self.lock:
                t=self.get(tid)
                if t.get('agent',{}).get('id')!=run_id:return
                a.update(status='paused' if isinstance(e,AgentPause) else 'blocked',message=str(e) if isinstance(e,(AgentPause,ModelError,AppError)) else '处理未完成，已保留进度，请重试。')
                a['pause_requested']=False;t.update(agent=a,busy=False,status='已暂停',error='' if isinstance(e,AgentPause) else a['message']);self.event(t,a['message']);self.save(t)

    def confirm_agent(self,tid,revision):
        # Explicit user review is the only route to final delivery; no model calls.
        with self.lock:
            t=self.get(tid);self.ensure_idle(t,revision)
            if t.get('agent',{}).get('status') in ('paused','blocked','waiting') and t['agent'].get('basis')==basis(t):raise AppError('Agent 还有未完成步骤，请继续处理或先手工完成任务书。')
            existing=next((e for e in reversed(t['exports']) if e['revision']==revision and e.get('agent_delivery')),None)
            if existing:return self.view(tid)
            self.review(tid,revision)
            kind='docx' if t['brief'].get('sample',{}).get('real') else 'both'
            result=self.export(tid,revision,kind);item=result['exports'][-1]
            for f in item['files']:
                if f['kind']=='docx':
                    doc=Document(f['path'])
                    if not doc.tables:raise AppError('生成文件缺少表格，请重新生成。')
                    body='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
                    expected=[t['brief'][k] for k in ('scenario','purpose','objective')]+[x['detail'] for x in t['brief']['outputs']]
                    if any(value not in body for value in expected):raise AppError('文件内容与已确认任务书不一致，请重新生成。')
                    for p in doc.element.body.xpath('.//w:p'):
                        pr=p.get_or_add_pPr();style=pr.find(qn('w:pStyle'))
                        if style is not None and style.get(qn('w:val')) in ('Title','Heading1','Heading2','Heading3'):continue
                        marker=pr.find(qn('w:keepNext'))
                        if marker is None or marker.get(qn('w:val'))!='0':raise AppError('正文格式检查未通过，请重新生成。')
                else:
                    wb=load_workbook(f['path'],read_only=True)
                    try:
                        if wb.active.max_column!=4+t['brief'].get('sample',{}).get('count',3):raise AppError('输入样例表列数检查未通过。')
                    finally:wb.close()
            t=self.get(tid);t['exports'][-1].update(agent_delivery=True,structural_check='passed',visual_review='pending')
            if t.get('agent'):
                if t['agent'].get('basis')!=basis(t):t['agent']['changes']=[]
                t['agent'].update(status='done',phase='done',basis=basis(t),message='已生成交付文件，内容结构检查通过；请打开文件核对排版。')
            self.save(t);return self.view(tid)
